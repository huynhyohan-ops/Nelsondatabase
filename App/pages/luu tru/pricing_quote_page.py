# pricing_quote_page.py
# Giữ nguyên cấu trúc tổng thể; chỉ thay đổi UI/mapping theo yêu cầu

import streamlit as st
import pandas as pd
import os
from datetime import date, datetime
from pathlib import Path
from docx import Document

import openpyxl
import re

from common.cost_engine import generate_quote
from common.models import (
    CustomerInfo,
    ShipmentInfo,
    ContainerPlanItem,
    EngineOptions,
    QuoteRequest,
    load_master,
)
from menu import top_menu
from common.style import inject_global_css


# ========================== UTILS (FORMAT/MAPPING ONLY) ==========================
def _fmt_dmmm(dt_value) -> str:
    """
    Trả về 'DMMM' viết tắt tháng IN HOA, không 0 dẫn.
    Hỗ trợ đầu vào dạng datetime/date/str; lỗi => '-'
    Ví dụ: 2025-12-04 -> '4DEC'
    """
    if pd.isna(dt_value) or dt_value in ("", None):
        return "-"
    try:
        if isinstance(dt_value, (pd.Timestamp, datetime, date)):
            d = pd.to_datetime(dt_value)
        else:
            d = pd.to_datetime(str(dt_value))
        return f"{int(d.day)}{d.strftime('%b').upper()}"
    except Exception:
        return "-"


def _fmt_valid(eff, exp) -> str:
    """
    Ghép 'DMMM - DMMM'. Nếu thiếu 1 đầu mốc => trả '-' (tránh gây hiểu nhầm).
    """
    s1, s2 = _fmt_dmmm(eff), _fmt_dmmm(exp)
    if s1 == "-" or s2 == "-":
        return "-"
    return f"{s1} - {s2}"


def _map_from_master(master_df: pd.DataFrame, key: dict, field: str):
    """
    Map 1 field từ master theo (POL, POD, PlaceOfDelivery, Carrier).
    Không thay đổi logic tính; chỉ lấy dữ liệu thô để hiển thị.
    """
    try:
        m = master_df[
            (master_df["POL"].astype(str).str.upper() == str(key["pol"]).upper())
            & (master_df["POD"].astype(str).str.upper() == str(key["pod"]).upper())
            & (master_df["PlaceOfDelivery"].astype(str).str.upper() == str(key["place"]).upper())
            & (master_df["Carrier"].astype(str).str.upper() == str(key["carrier"]).upper())
        ]
        if not m.empty and field in m.columns:
            return m.iloc[0][field]
    except Exception:
        pass
    return "-"


def _tt_compact(min_tt, max_tt) -> str:
    """
    'min-maxd' chuẩn hóa hiển thị thời gian transit. Nếu thiếu, fallback '-'.
    """
    try:
        if pd.isna(min_tt) or pd.isna(max_tt):
            return "-"
        return f"{int(min_tt)}-{int(max_tt)}d"
    except Exception:
        return "-"


# ========================== VERSION KPI ==========================
def get_latest_version_from_master(data_dir="Data") -> str:
    """
    Đọc tên sheet dạng DDMMMNO<VERSION> từ 'Master_FullPricing.xlsx'.
    Trả về sheet cuối phù hợp (không can thiệp tính toán).
    """
    master_path = os.path.join(data_dir, "Master_FullPricing.xlsx")
    if not os.path.exists(master_path):
        return "N/A"

    try:
        wb = openpyxl.load_workbook(master_path, read_only=True)
        sheets = wb.sheetnames

        versions = [s for s in sheets if re.match(r"^\d{1,2}[A-Z]{3}NO\d+$", s.upper())]
        if not versions:
            return "N/A"
        # Giữ nguyên nguyên tắc cũ: lấy sheet cuối cùng (không đổi thuật toán)
        return versions[-1].upper()
    except Exception:
        return "N/A"


# ========================== PIPELINE EXTRACTION ==========================
@st.cache_data(ttl=3600)
def extract_pipeline_data(doc_path: str) -> dict:
    data = {"default_markup": {}, "carriers": []}
    try:
        doc = Document(doc_path)
        for para in doc.paragraphs:
            t = para.text.strip()
            if not t:
                continue
            if t.startswith("CARRIER:"):
                name = t.split(":", 1)[1].strip().upper()
                data["carriers"].append(name)
            elif t.startswith("DEFAULT-MARKUP:"):
                _, v = t.split(":", 1)
                p = v.split("=")
                if len(p) == 2:
                    data["default_markup"][p[0].strip().upper()] = float(p[1])
    except Exception:
        pass
    return data


# ========================== KPI SUMMARY CARD ==========================
def render_summary_cards(master_df: pd.DataFrame):
    if master_df.empty:
        return
    total_rows = len(master_df)
    total_carriers = master_df["Carrier"].nunique()
    version = get_latest_version_from_master()

    st.markdown("### 📊 Summary KPI Card")
    col1, col2, col3 = st.columns(3)

    kpi_style = """
        <style>
        .kpi-card {
            background-color: #f6f8fa;
            padding: 10px;
            border-radius: 8px;
            text-align: center;
            font-weight: bold;
            height: 80px;
            display: flex;
            justify-content: center;
            align-items: center;
            font-size: 18px;
        }
        </style>
    """
    st.markdown(kpi_style, unsafe_allow_html=True)

    with col1:
        st.markdown(f"<div class='kpi-card'>Master Rows<br>{total_rows}</div>", unsafe_allow_html=True)

    with col2:
        st.markdown(f"<div class='kpi-card'>Unique Carriers<br>{total_carriers}</div>", unsafe_allow_html=True)

    with col3:
        # Hiển thị đúng format sheet, ví dụ 12DECNO1
        st.markdown(f"<div class='kpi-card'>Version<br>{version}</div>", unsafe_allow_html=True)


# ========================== MAIN PAGE ==========================
def render_quote_page():
    top_menu(active="pricing")
    inject_global_css()
    st.markdown("### ⚓ Ocean Freight Quotation")
    st.caption("Live Preview Cost + Ranking by Container Type")
    st.markdown("---")

    # --- Load master ---
    master_df = load_master()
    if master_df.empty:
        st.error("Không tìm thấy dữ liệu MasterFullPricing.")
        return

    # --- Load pipeline ---
    pipeline_path = Path("Pipeline.docx") if Path("Pipeline.docx").exists() else Path("data/Pipeline.docx")
    pipeline_data = extract_pipeline_data(str(pipeline_path))

    # --- KPI Summary Cards ---
    render_summary_cards(master_df)
    st.markdown("---")

    # ========== SHIPMENT INFORMATION ==========
    st.subheader("🚢 Shipment Information")

    pol_list = sorted(master_df["POL"].dropna().astype(str).str.upper().unique().tolist())
    pol_selected = st.segmented_control("Select POL", options=pol_list, selection_mode="single")

    place_list = sorted(
        master_df.loc[master_df["POL"].astype(str).str.upper() == str(pol_selected).upper(), "PlaceOfDelivery"]
        .dropna().astype(str).str.upper().unique().tolist()
    )
    # (1) Cho phép chọn NHIỀU Place Of Delivery (chỉ thay đổi UI)
    places_selected = st.multiselect(
        "Place of Delivery",
        options=place_list,
        default=place_list[:1] if place_list else [],
        key="multi_place_of_delivery",
    )

    # POD phụ thuộc vào các Place đã chọn (union)
    if places_selected:
        pod_list = sorted(
            master_df.loc[
                master_df["PlaceOfDelivery"].astype(str).str.upper().isin([p.upper() for p in places_selected]),
                "POD",
            ].dropna().astype(str).str.upper().unique().tolist()
        )
    else:
        pod_list = []
    pod_selected = st.selectbox("POD (optional)", [""] + pod_list if pod_list else [""])

    col1, col2 = st.columns(2)
    with col1:
        fak_reefer = st.toggle("FAK / REEFER", value=False)
    with col2:
        soc = st.toggle("SOC", value=False)

    st.markdown("---")

    # ========== CONTAINER PLAN ==========
    st.subheader("📦 Container Plan")
    container_options = ["20GP", "40GP", "40HQ", "45", "40NOR"] if not fak_reefer else ["20RF", "40RF"]

    container_selected = st.multiselect(
        "Select Container Types",
        options=container_options,
        default=["40HQ"] if "40HQ" in container_options else [container_options[0]],
        key="multi_containers"
    )
    st.caption(f"Selected Containers: {', '.join(container_selected) if container_selected else '-'}")
    st.markdown("---")

    # ========== 💲 LIVE PRICING PREVIEW ==========
    st.subheader("💲 Live Pricing Preview")

    if not places_selected or not container_selected:
        st.info("Vui lòng chọn ít nhất 1 Place Of Delivery và 1 Container.")
        return

    cust = CustomerInfo(name="Preview", email="-")

    opts = EngineOptions(currency="USD", max_options_per_quote=10)
    opts.markup_map = pipeline_data.get("default_markup", {})

    # Gom preview cho N place_of_delivery (không thay đổi thuật toán tính)
    preview_frames = []
    for place in places_selected:
        shipment = ShipmentInfo(
            pol=pol_selected,
            place_of_delivery=place,
            pod=pod_selected,
            cargo_ready_date=date.today().isoformat(),
            commodity_type="REEFER" if fak_reefer else "FAK",
            is_soc=soc,
        )
        containers = [ContainerPlanItem(type=c, quantity=1) for c in container_selected]

        preview_req = QuoteRequest(
            customer=cust,
            shipment=shipment,
            containers=containers,
            engine_options=opts,
        )
        res = generate_quote(master_df, preview_req)
        if "options" in res:
            df_part = pd.DataFrame(res["options"])
            if not df_part.empty:
                # gắn lại place để tránh mất ngữ cảnh khi concat
                df_part["place_of_delivery"] = place
                preview_frames.append(df_part)

    if not preview_frames:
        st.warning("Không có dữ liệu phù hợp filter hiện tại.")
        return

    df_preview = pd.concat(preview_frames, ignore_index=True)

    # Chọn container để xếp hạng (giữ nguyên tiêu chí cũ)
    if len(container_selected) > 1:
        if "40HQ" in container_selected:
            container_rank = "40HQ"
        elif "40GP" in container_selected:
            container_rank = "40GP"
        else:
            container_rank = container_selected[-1]
    else:
        container_rank = container_selected[0]

    # ---------- XÂY BẢNG HIỂN THỊ + VALID ----------
    display_data = pd.DataFrame({
        "POL": df_preview["pol"],
        "POD": df_preview["pod"],
        "PLACE OF DELIVERY": df_preview["place_of_delivery"],
        "CARRIER": df_preview["carrier"],
        "SERVICE": df_preview["service"],
        "VESSEL": df_preview["vessel"],
        "ETD-ETA": df_preview.apply(lambda r: f"{r['etd']} → {r['eta']}", axis=1),
        "TRANSIT TIME": df_preview.apply(lambda r: _tt_compact(r.get("transit_min"), r.get("transit_max")), axis=1),
    })

    # Map VALID từ master theo khóa (POL,POD,Place,Carrier)
    valid_list = []
    for _, r in df_preview.iterrows():
        key = {
            "pol": r.get("pol", ""),
            "pod": r.get("pod", ""),
            "place": r.get("place_of_delivery", ""),
            "carrier": r.get("carrier", ""),
        }
        eff = _map_from_master(master_df, key, "EffectiveDate")
        exp = _map_from_master(master_df, key, "ExpirationDate")
        valid_list.append(_fmt_valid(eff, exp))
    display_data["VALID"] = valid_list

    # Mapping bổ sung giữ nguyên (không ảnh hưởng tính)
    if "RoutingNote" in df_preview.columns:
        display_data["RoutingNote"] = df_preview["RoutingNote"].values
    if "CommodityType" in df_preview.columns:
        display_data["CommodityType"] = df_preview["CommodityType"].values

    # Thêm các cột container price
    for c in container_selected:
        display_data[c] = df_preview["container_rates"].apply(lambda x: x.get(c, "-"))

    # Ranking giữ nguyên logic cũ
    display_data["RANK TARGET"] = df_preview["container_rates"].apply(lambda x: x.get(container_rank, None))
    display_data = display_data[display_data["RANK TARGET"].notna()]
    display_data = display_data.sort_values(by="RANK TARGET", ascending=True).head(10)
    display_data["RANK"] = range(1, len(display_data) + 1)

    ordered_cols = [
        "RANK",
        "POL",
        "POD",
        "PLACE OF DELIVERY",
        "CARRIER",
        "SERVICE",
        "VESSEL",
        "ETD-ETA",
        "VALID",          # <== cột mới theo yêu cầu
        "TRANSIT TIME",   # giữ dạng 'min-maxd'
        "RoutingNote",
        "CommodityType",
    ] + container_selected
    # Lọc những cột có thật để không lỗi trong case thiếu
    ordered_cols = [c for c in ordered_cols if c in display_data.columns]

    display_data = display_data[ordered_cols]
    st.markdown(f"**📈 Top 10 best rates ranked by `{container_rank}` container**")
    st.dataframe(display_data, use_container_width=True)

    # ========== 📈 MARKUP SETTINGS ==========
    st.markdown("---")
    with st.expander("📈 Carrier Markup Settings (Optional)", expanded=False):
        st.markdown("Điều chỉnh giá từng hãng tàu")
        markup_map = {}
        carriers = sorted(set(display_data["CARRIER"].dropna().astype(str).tolist()))
        for c in carriers:
            base = int(pipeline_data.get("default_markup", {}).get(c, 50))
            markup_map[c] = st.slider(f"{c} Markup (USD)", 0, 200, base, step=5)
        st.json(markup_map)

    # ========== 🚀 GENERATE QUOTE ==========
    st.markdown("---")
    st.subheader("🚀 Generate Quote")

    colA, colB = st.columns([0.4, 0.6])
    with colA:
        customer_name = st.text_input("Customer Name", "Demo Customer")
        email = st.text_input("Email", "")
    with colB:
        cargo_ready_date = st.date_input("Cargo Ready Date", date.today())

    if st.button("Generate Quote", use_container_width=True):
        st.info("Đang tạo báo giá chính thức...")
        try:
            opts.markup_map = markup_map
            cust = CustomerInfo(name=customer_name, email=email)

            final_frames = []
            for place in places_selected:
                shipment = ShipmentInfo(
                    pol=pol_selected,
                    place_of_delivery=place,
                    pod=pod_selected,
                    cargo_ready_date=cargo_ready_date.isoformat(),
                    commodity_type="REEFER" if fak_reefer else "FAK",
                    is_soc=soc,
                )
                containers = [ContainerPlanItem(type=c, quantity=1) for c in container_selected]

                req = QuoteRequest(
                    customer=cust,
                    shipment=shipment,
                    containers=containers,
                    engine_options=opts,
                )
                result = generate_quote(master_df, req)
                if "options" in result:
                    df_part = pd.DataFrame(result["options"])
                    if not df_part.empty:
                        df_part["place_of_delivery"] = place
                        final_frames.append(df_part)

            if not final_frames:
                st.warning("Không có dữ liệu để hiển thị báo giá.")
                return

            df_result = pd.concat(final_frames, ignore_index=True)
            st.success("✅ Quote generated successfully.")

            # Mapping và build bảng Final (giữ logic; thêm VALID)
            df_display = pd.DataFrame({
                "POL": df_result["pol"],
                "POD": df_result["pod"],
                "PLACE OF DELIVERY": df_result["place_of_delivery"],
                "CARRIER": df_result["carrier"],
                "SERVICE": df_result["service"],
                "VESSEL": df_result["vessel"],
                "ETD-ETA": df_result.apply(lambda r: f"{r['etd']} → {r['eta']}", axis=1),
                "TRANSIT TIME": df_result.apply(lambda r: _tt_compact(r.get("transit_min"), r.get("transit_max")), axis=1),
            })

            valid_list_final = []
            for _, r in df_result.iterrows():
                key = {
                    "pol": r.get("pol", ""),
                    "pod": r.get("pod", ""),
                    "place": r.get("place_of_delivery", ""),
                    "carrier": r.get("carrier", ""),
                }
                eff = _map_from_master(master_df, key, "EffectiveDate")
                exp = _map_from_master(master_df, key, "ExpirationDate")
                valid_list_final.append(_fmt_valid(eff, exp))
            df_display["VALID"] = valid_list_final

            # Thêm Routing/Commodity như trước (nếu có)
            if "RoutingNote" in df_result.columns:
                df_display["RoutingNote"] = df_result["RoutingNote"].values
            if "CommodityType" in df_result.columns:
                df_display["CommodityType"] = df_result["CommodityType"].values

            for c in container_selected:
                df_display[c] = df_result["container_rates"].apply(lambda x: x.get(c, "-"))

            if len(container_selected) > 1:
                if "40HQ" in container_selected:
                    container_rank = "40HQ"
                elif "40GP" in container_selected:
                    container_rank = "40GP"
                else:
                    container_rank = container_selected[-1]
            else:
                container_rank = container_selected[0]

            df_display["RANK TARGET"] = df_result["container_rates"].apply(lambda x: x.get(container_rank, None))
            df_display = df_display[df_display["RANK TARGET"].notna()]
            df_display = df_display.sort_values(by="RANK TARGET", ascending=True).head(10)
            df_display["RANK"] = range(1, len(df_display) + 1)

            ordered_cols = [
                "RANK", "POL", "POD", "PLACE OF DELIVERY", "CARRIER", "SERVICE",
                "VESSEL", "ETD-ETA", "VALID", "TRANSIT TIME", "RoutingNote", "CommodityType"
            ] + container_selected
            ordered_cols = [c for c in ordered_cols if c in df_display.columns]
            df_display = df_display[ordered_cols]

            st.markdown("### 📦 Final Quotation Preview (Ranked)")
            st.dataframe(df_display, use_container_width=True)

        except Exception as e:
            st.error(f"Lỗi khi tạo báo giá: {e}")


# ========================== ENTRY POINT ==========================
if __name__ == "__main__":
    render_quote_page()
